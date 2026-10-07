"""Server-resolved source previews; no client-provided filesystem paths."""
import io
from pathlib import Path
import uuid

import pymupdf
from fastapi import HTTPException
from sqlalchemy import select
from app.core.config import settings
from app.models.source_version import SourceVersion
from app.schemas.source_preview import SourceLocation
from app.services.media_tools import MediaTools

MIME = {'pdf':'application/pdf','pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation','video':'video/mp4'}


def contained_file(value: str, root: Path) -> Path:
    try:
        path = Path(value).resolve(strict=True)
        if not path.is_relative_to(root.resolve()) or not path.is_file():
            raise HTTPException(404,'Source asset is unavailable')
        return path
    except (OSError,ValueError):
        raise HTTPException(404,'Source asset is unavailable') from None


async def resolve_source(db, course_id: uuid.UUID, source_id: uuid.UUID):
    source = await db.scalar(select(SourceVersion).where(SourceVersion.id==source_id,SourceVersion.course_id==course_id))
    if source is None: raise HTTPException(404,'Source not found in this course')
    kind = source.source_type.value
    if kind not in MIME: raise HTTPException(415,'Unsupported source type')
    path = contained_file(source.storage_path,Path(settings.upload_dir)/str(course_id))
    return source,path


def validate_location(path: Path, kind: str, location: SourceLocation) -> dict:
    if location.kind != {'pdf':'page','pptx':'slide','video':'time'}.get(kind):
        raise HTTPException(422,'Locator kind does not match this source')
    try:
        if kind=='pdf':
            with pymupdf.open(path) as doc: count=len(doc)
        elif kind=='pptx':
            from pptx import Presentation
            count=len(Presentation(path).slides)
        else:
            duration=float(MediaTools(timeout_seconds=30).probe_media(str(path))['format']['duration'])
            if location.start > duration or (location.end is not None and location.end>duration):
                raise HTTPException(422,'Timestamp exceeds source duration')
            return dict(duration=duration)
        if location.number>count: raise HTTPException(422,'Page or slide exceeds source length')
        return dict(count=count)
    except HTTPException: raise
    except Exception:
        raise HTTPException(422,'Source cannot be previewed') from None


def pdf_page(path: Path, number: int) -> bytes:
    with pymupdf.open(path) as doc:
        page=doc[number-1]
        scale=min(2.,1800/max(page.rect.width,page.rect.height))
        return page.get_pixmap(matrix=pymupdf.Matrix(scale,scale),alpha=False).tobytes('png')


def slide_preview(path: Path, number: int) -> tuple[bytes,list[str]]:
    """Render source-owned text, pictures, simple shapes and tables at their slide coordinates.

    Complex Office features are explicitly marked, never replaced by invented content.
    The original presentation remains downloadable. No scripts/external assets execute.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from PIL import Image
    prs=Presentation(path); slide=prs.slides[number-1]; scale=72/914400
    warnings=[]
    document=pymupdf.open()
    page=document.new_page(width=prs.slide_width*scale,height=prs.slide_height*scale)
    def shapes(items):
        for shape in items:
            box=pymupdf.Rect(shape.left*scale,shape.top*scale,(shape.left+shape.width)*scale,(shape.top+shape.height)*scale)
            if shape.shape_type==MSO_SHAPE_TYPE.LINE:
                page.draw_line(box.top_left,box.bottom_right,color=(.2,.2,.2))
                warnings.append('Connector styling is simplified')
                continue
            if box.is_empty: continue
            if shape.shape_type==MSO_SHAPE_TYPE.GROUP:
                warnings.append('Grouped shape transforms require the original presentation'); continue
            if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                try:
                    with Image.open(io.BytesIO(shape.image.blob)) as image:
                        output=io.BytesIO(); image.convert('RGB').save(output,format='PNG')
                    page.insert_image(box,stream=output.getvalue())
                except Exception: warnings.append('An embedded image could not be rendered')
            elif getattr(shape,'has_table',False):
                table=shape.table
                x=box.x0
                for column_index,column in enumerate(table.columns):
                    y=box.y0
                    for row_index,row in enumerate(table.rows):
                        cell=pymupdf.Rect(x,y,x+column.width*scale,y+row.height*scale)
                        page.draw_rect(cell,color=(.4,.4,.4),width=.5)
                        page.insert_textbox(cell+ (2,2,-2,-2),table.cell(row_index,column_index).text,fontsize=9)
                        y+=row.height*scale
                    x+=column.width*scale
            elif getattr(shape,'has_text_frame',False):
                # Text is taken only from the original slide. Rectangles/ellipses preserve useful layout.
                if shape.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
                    auto=str(shape.auto_shape_type)
                    if 'OVAL' in auto: page.draw_oval(box,color=(.3,.3,.3))
                    elif 'RECTANGLE' in auto: page.draw_rect(box,color=(.3,.3,.3))
                    else: warnings.append('Some shape outlines are simplified')
                text=shape.text_frame.text
                size=14
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size: size=min(40,max(6,run.font.size.pt)); break
                while size>=6:
                    shape_draw=page.new_shape()
                    if shape_draw.insert_textbox(box+(3,3,-3,-3),text,fontsize=size)>=0:
                        shape_draw.commit(); break
                    size-=1
                else: warnings.append('Some text does not fit; inspect the original presentation')
            else:
                warnings.append('An Office object requires the original presentation')
    shapes(slide.shapes)
    image=page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False).tobytes('png')
    document.close()
    return image,list(dict.fromkeys(warnings))
