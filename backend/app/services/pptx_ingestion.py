import logging
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.evidence_unit import EvidenceUnit, ExtractionStatus, Modality
from app.models.source_version import ProcessingStatus, SourceType, SourceVersion
from app.services.pdf_ingestion import (
    AlreadyProcessedError,
    ProcessingResult,
    _count_existing_evidence,
)
from app.services.derived_storage import derived_dir_for, safe_derived_filename

logger = logging.getLogger(__name__)


class PPTXProcessingError(Exception):
    """Raised when PPTX processing fails unrecoverably."""


class NotAPPTXError(Exception):
    """Raised when processing is requested for a non-PPTX source."""


async def process_pptx(
    db: AsyncSession,
    *,
    source_version: SourceVersion,
) -> ProcessingResult:
    """Process a PPTX SourceVersion: extract text shapes, tables, and images.

    Raises:
        NotAPPTXError: if source_type is not PPTX
        AlreadyProcessedError: if already in ``ready`` status (carries result)

    Returns:
        ProcessingResult with counts and final status.
    """
    sv_id = source_version.id

    # guard: non-PPTX
    if source_version.source_type != SourceType.pptx:
        raise NotAPPTXError(
            f"SourceVersion {sv_id} is {source_version.source_type.value}, not PPTX"
        )

    # guard: already processed
    if source_version.processing_status == ProcessingStatus.ready:
        existing_count = await _count_existing_evidence(db, sv_id)

        slide_count = None
        pptx_path = Path(source_version.storage_path)
        if pptx_path.exists():
            try:
                prs = Presentation(str(pptx_path))
                slide_count = len(prs.slides)
            except Exception:
                pass

        result = ProcessingResult(
            source_version_id=sv_id,
            processing_status="ready",
            already_processed=True,
            page_count=None,
            slide_count=slide_count,
            **existing_count,
        )
        raise AlreadyProcessedError(result)

    # transition: uploaded/failed → processing
    source_version.processing_status = ProcessingStatus.processing
    await db.commit()
    await db.refresh(source_version)

    logger.info(
        "PPTX processing started: source_version_id=%s, file=%s",
        sv_id,
        source_version.original_filename,
    )

    try:
        result = await _extract_pptx(db, source_version=source_version)
    except Exception as exc:
        # transition: processing → failed
        await db.rollback()
        logger.error("PPTX processing failed: source_version_id=%s, error=%s", sv_id, exc)
        source_version.processing_status = ProcessingStatus.failed
        await db.commit()
        raise PPTXProcessingError(str(exc)) from exc

    # transition: processing → ready
    source_version.processing_status = ProcessingStatus.ready
    await db.commit()

    logger.info(
        "PPTX processing completed: source_version_id=%s, slides=%d, "
        "text=%d, tables=%d, images=%d",
        sv_id,
        result.slide_count or 0,
        result.text_evidence_count,
        result.table_evidence_count,
        result.image_evidence_count,
    )

    return result


def _table_to_markdown(table) -> str:
    """Convert a pptx Table into a Markdown table."""
    if not table.rows:
        return ""

    def _cell(text: str | None) -> str:
        if text is None:
            return ""
        # Preserve multi-line by replacing newline with <br>
        # Escape pipe characters to preserve markdown table
        return str(text).replace("|", "\\|").replace("\n", "<br>").strip()

    lines: list[str] = []

    # Process header
    header = [_cell(cell.text_frame.text if getattr(cell, "text_frame", None) else "") for cell in table.rows[0].cells]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join("---" for _ in header) + " |")

    # Process body
    for row in list(table.rows)[1:]:
        cells = [_cell(cell.text_frame.text if getattr(cell, "text_frame", None) else "") for cell in row.cells]
        while len(cells) < len(header):
            cells.append("")
        lines.append("| " + " | ".join(cells) + " |")

    return "\n".join(lines)


async def _extract_pptx(
    db: AsyncSession,
    *,
    source_version: SourceVersion,
) -> ProcessingResult:
    """Run extraction passes on the PPTX and persist EvidenceUnits."""
    sv_id = source_version.id
    course_id = source_version.course_id
    pptx_path = Path(source_version.storage_path)

    if not pptx_path.exists():
        raise FileNotFoundError(f"PPTX file not found: {pptx_path}")

    derived_dir = derived_dir_for(sv_id)

    result = ProcessingResult(
        source_version_id=sv_id,
        processing_status="ready",
    )
    result.text_evidence_count = 0
    result.table_evidence_count = 0
    result.image_evidence_count = 0

    evidence_units: list[EvidenceUnit] = []

    try:
        prs = Presentation(str(pptx_path))
    except Exception as exc:
        raise PPTXProcessingError(f"Cannot open PPTX: {exc}") from exc

    slide_width = prs.slide_width
    slide_height = prs.slide_height
    result.slide_count = len(prs.slides)
    result.page_count = None

    for slide_index, slide in enumerate(prs.slides):
        slide_number = slide_index + 1

        # Track image count per slide for deterministic artifacts
        slide_image_index = 0

        for shape_index, shape in enumerate(slide.shapes):
            left = shape.left
            top = shape.top
            width = shape.width
            height = shape.height

            # Avoid missing coordinates crashing extraction
            bbox = None
            if left is not None and top is not None and width is not None and height is not None:
                bbox = [left, top, left + width, top + height]

            location_metadata = {
                "slide_index": slide_index,
                "slide_number": slide_number,
                "shape_id": shape.shape_id if hasattr(shape, "shape_id") else None,
                "shape_name": shape.name if hasattr(shape, "name") else None,
                "shape_index": shape_index,
                "bounding_box": bbox,
                "coordinate_unit": "emu",
                "slide_width": slide_width,
                "slide_height": slide_height,
            }

            # Table checking
            if getattr(shape, "has_table", False):
                md_text = _table_to_markdown(shape.table)
                if md_text.strip():
                    loc_meta = location_metadata.copy()
                    loc_meta["element_type"] = "table"
                    loc_meta["table_rows"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    eu = EvidenceUnit(
                        course_id=course_id,
                        source_version_id=sv_id,
                        modality=Modality.table,
                        original_text=md_text,
                        extracted_description=None,
                        topic=None,
                        subtopic=None,
                        extraction_status=ExtractionStatus.extracted,
                        location_metadata=loc_meta,
                    )
                    evidence_units.append(eu)
                    result.table_evidence_count += 1
                continue

            # Picture checking
            shape_type = getattr(shape, "shape_type", None)
            if shape_type == MSO_SHAPE_TYPE.PICTURE and hasattr(shape, "image"):
                image_bytes = shape.image.blob
                ext = getattr(shape.image, "ext", "png")

                filename = safe_derived_filename(
                    prefix="pptx_img", page_index=slide_index, img_index=slide_image_index, ext=ext,
                )
                image_path = derived_dir / filename
                image_path.write_bytes(image_bytes)

                slide_image_index += 1

                img_size = getattr(shape.image, "size", None)

                loc_meta = location_metadata.copy()
                loc_meta["element_type"] = "image"
                loc_meta["derived_path"] = str(image_path)
                loc_meta["image_extension"] = ext
                if img_size and len(img_size) == 2:
                    loc_meta["image_width"] = img_size[0]
                    loc_meta["image_height"] = img_size[1]

                eu = EvidenceUnit(
                    course_id=course_id,
                    source_version_id=sv_id,
                    modality=Modality.image,
                    original_text=None,
                    extracted_description=None,
                    topic=None,
                    subtopic=None,
                    extraction_status=ExtractionStatus.extracted,
                    location_metadata=loc_meta,
                )
                evidence_units.append(eu)
                result.image_evidence_count += 1
                continue

            # Text checking
            if getattr(shape, "has_text_frame", False):
                raw_text = shape.text
                if raw_text.strip():
                    loc_meta = location_metadata.copy()
                    loc_meta["element_type"] = "text"
                    eu = EvidenceUnit(
                        course_id=course_id,
                        source_version_id=sv_id,
                        modality=Modality.text,
                        original_text=raw_text,
                        extracted_description=None,
                        topic=None,
                        subtopic=None,
                        extraction_status=ExtractionStatus.extracted,
                        location_metadata=loc_meta,
                    )
                    evidence_units.append(eu)
                    result.text_evidence_count += 1

    if evidence_units:
        db.add_all(evidence_units)
        await db.flush()

    return result
