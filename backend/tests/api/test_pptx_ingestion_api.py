import os
import uuid
import hashlib
from pathlib import Path

import pytest
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession
from pptx import Presentation
from pptx.util import Inches

from app.models.evidence_unit import EvidenceUnit
from app.models.source_version import SourceVersion


@pytest.fixture
def tmp_pptx(tmp_path) -> Path:
    """Fixture that generates a programmatic PPTX for testing."""
    prs = Presentation()

    # Slide 1: Normal text
    blank_slide_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_slide_layout)
    txBox1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    tf1 = txBox1.text_frame
    tf1.text = "  \n  Slide 1 Text  \n  "

    # Add empty text shape (should be skipped)
    txBox_empty = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
    tf_empty = txBox_empty.text_frame
    tf_empty.text = "   \n  "

    # Slide 2: Unicode text
    slide2 = prs.slides.add_slide(blank_slide_layout)
    txBox2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    tf2 = txBox2.text_frame
    tf2.text = "Unicode 🌟 ñ 測試"

    # Slide 3: Table (including blank cell and literal |)
    slide3 = prs.slides.add_slide(blank_slide_layout)
    table_shape = slide3.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Process"
    table.cell(0, 1).text = "Arrival"
    table.cell(0, 2).text = "Burst"
    table.cell(1, 0).text = "P1"
    # Blank cells in numeric data
    table.cell(1, 1).text = ""
    table.cell(1, 2).text = "8"
    table.cell(2, 0).text = "P2|Pipe"
    table.cell(2, 1).text = "5"
    table.cell(2, 2).text = "4"

    # Adding a separate text box on the table slide to test mixed content
    txBox3 = slide3.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1))
    txBox3.text_frame.text = "Text with table"

    # Slide 4: Image
    slide4 = prs.slides.add_slide(blank_slide_layout)
    dummy_img_path = tmp_path / "dummy.png"
    dummy_img_path.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
        b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00"
        b"\x01\n\x00\x02\x00\x00\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    pic = slide4.shapes.add_picture(str(dummy_img_path), Inches(1), Inches(1))

    # Also test an empty table (should be skipped) - python-pptx doesn't let us create a 0-row table easily
    # We will simulate missing logic or empty table behavior on text empty check.

    pptx_file = tmp_path / "test.pptx"
    prs.save(str(pptx_file))
    return pptx_file


@pytest.fixture
def corrupt_pptx(tmp_path) -> Path:
    """Fixture that generates a corrupt PPTX for testing."""
    pptx_file = tmp_path / "corrupt.pptx"
    pptx_file.write_bytes(b"PK\x03\x04 not a real pptx file content")
    return pptx_file


def _sha256(path: Path) -> str:
    hasher = hashlib.sha256()
    hasher.update(path.read_bytes())
    return hasher.hexdigest()


@pytest.mark.asyncio
async def test_process_pptx_all_evidence(client: AsyncClient, tmp_pptx: Path, db_session: AsyncSession):
    # Calculate initial hash to ensure immutability
    original_hash = _sha256(tmp_pptx)

    # 1. Create course
    resp = await client.post("/courses", json={"name": "Process PPTX Dispatcher"})
    assert resp.status_code == 201
    course_id = resp.json()["id"]

    # 2. Upload PPTX
    resp = await client.post(
        f"/courses/{course_id}/sources",
        files={"file": ("test.pptx", tmp_pptx.read_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert resp.status_code == 201
    source_version_id = resp.json()["id"]
    storage_path = Path(resp.json()["storage_path"])

    upload_hash = _sha256(storage_path)
    assert original_hash == upload_hash

    # 3. Trigger processing
    resp = await client.post(f"/courses/{course_id}/sources/{source_version_id}/process")
    assert resp.status_code == 201
    data = resp.json()
    assert data["processing_status"] == "ready"
    assert data["slide_count"] == 4
    assert data["text_evidence_count"] == 3  # Slide 1 Text, Unicode text, Text with table
    assert data["table_evidence_count"] == 1
    assert data["image_evidence_count"] == 1

    # Immutability check
    post_process_hash = _sha256(storage_path)
    assert original_hash == post_process_hash

    # 4. Verify evidence in DB
    from sqlalchemy import select

    result = await db_session.execute(
        select(EvidenceUnit)
        .where(EvidenceUnit.source_version_id == uuid.UUID(source_version_id))
    )
    evidence_units = result.scalars().all()
    assert len(evidence_units) == 5

    texts = [eu for eu in evidence_units if eu.modality.value == "text"]
    tables = [eu for eu in evidence_units if eu.modality.value == "table"]
    images = [eu for eu in evidence_units if eu.modality.value == "image"]

    assert len(texts) == 3
    assert len(tables) == 1
    assert len(images) == 1

    # Table verification
    table_eu = tables[0]
    expected_md = (
        "| Process | Arrival | Burst |\n"
        "| --- | --- | --- |\n"
        "| P1 |  | 8 |\n"
        "| P2\\|Pipe | 5 | 4 |"
    )
    assert table_eu.original_text == expected_md
    assert table_eu.location_metadata["element_type"] == "table"
    assert table_eu.location_metadata["slide_index"] == 2
    assert table_eu.location_metadata["slide_number"] == 3
    assert table_eu.extracted_description is None

    # No duplicate text evidence generated from table
    for eu in texts:
        assert "Process" not in str(eu.original_text)
        assert "Burst" not in str(eu.original_text)

    # Image verification
    image_eu = images[0]
    assert image_eu.original_text is None
    assert image_eu.extracted_description is None
    assert image_eu.location_metadata["element_type"] == "image"
    assert "derived_path" in image_eu.location_metadata

    derived_path = Path(image_eu.location_metadata["derived_path"])
    assert derived_path.exists()
    assert derived_path.is_file()
    assert str(derived_path).endswith(".png")
    # Derived path should be outside upload
    assert derived_path.parent.parent.name == "derived"

    # IDEMPOTENCY check

    # We must count original derived files
    derived_files_count = len(list(derived_path.parent.glob("*")))
    assert derived_files_count == 1

    # Process again
    resp = await client.post(f"/courses/{course_id}/sources/{source_version_id}/process")
    assert resp.status_code == 200
    data2 = resp.json()
    assert data2["already_processed"] is True
    assert data2["text_evidence_count"] == 3
    assert data2["table_evidence_count"] == 1
    assert data2["image_evidence_count"] == 1

    # Re-verify DB and derived dir count
    result = await db_session.execute(
        select(EvidenceUnit)
        .where(EvidenceUnit.source_version_id == uuid.UUID(source_version_id))
    )
    assert len(result.scalars().all()) == 5
    assert len(list(derived_path.parent.glob("*"))) == 1


@pytest.mark.asyncio
async def test_process_corrupt_pptx(client: AsyncClient, corrupt_pptx: Path, db_session: AsyncSession):
    # Tests failure safety
    resp = await client.post("/courses", json={"name": "Corrupt PPTX"})
    assert resp.status_code == 201
    course_id = resp.json()["id"]

    resp = await client.post(
        f"/courses/{course_id}/sources",
        files={"file": ("corrupt.pptx", corrupt_pptx.read_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert resp.status_code == 201
    source_version_id = resp.json()["id"]

    # Trigger processing
    resp = await client.post(f"/courses/{course_id}/sources/{source_version_id}/process")
    assert resp.status_code == 422
    assert "PPTX processing failed" in resp.json()["detail"]

    # DB verify no partial evidence exists
    from sqlalchemy import select
    result = await db_session.execute(
        select(EvidenceUnit)
        .where(EvidenceUnit.source_version_id == uuid.UUID(source_version_id))
    )
    assert len(result.scalars().all()) == 0
