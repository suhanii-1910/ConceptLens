"""Real file previews and source/visual ownership; providers are explicitly fixtures."""
import json
import uuid
from pathlib import Path
import pymupdf
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from app.core.config import settings
from app.models import Course, SourceVersion, EvidenceUnit
from app.services.visual_evidence import interpret, visual_asset
from app.services.source_preview import contained_file
from fastapi import HTTPException

async def source(db, kind='pdf'):
    course=Course(id=uuid.uuid4(),name='Source previews')
    root=Path(settings.upload_dir)/str(course.id);root.mkdir()
    path=root/('source.'+('mp4' if kind=='video' else kind))
    if kind=='pdf':
        doc=pymupdf.open()
        for text in ['First page','Second page']:
            page=doc.new_page();page.insert_text((50,50),text)
        doc.save(path);doc.close()
    elif kind=='pptx':
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1)).text='Original slide content';prs.save(path)
    else:path.write_bytes(b'fixture video')
    row=SourceVersion(id=uuid.uuid4(),course_id=course.id,original_filename='source.'+kind,file_hash='a'*64,
        mime_type='application/pdf',source_type=kind,storage_path=str(path),processing_status='ready')
    db.add_all([course,row]);await db.commit()
    return course,row,path

@pytest.mark.asyncio
@pytest.mark.parametrize('kind,locator', [('pdf','kind=page&number=2'),('pptx','kind=slide&number=1')])
async def test_actual_preview_and_original(db_session,client,kind,locator):
    c,s,path=await source(db_session,kind);base=f'/courses/{c.id}/sources/{s.id}'
    result=await client.get(base+'/location?'+locator);assert result.status_code==200
    assert result.json()['count']==(2 if kind=='pdf' else 1)
    preview=await client.get(result.json()['preview_url']);assert preview.content.startswith(b'\x89PNG')
    original=await client.get(base+'/file');assert original.content==path.read_bytes()
    assert original.headers['x-content-type-options']=='nosniff'
    assert 'source.' in original.headers['content-disposition']
    other=await client.get(f'/courses/{uuid.uuid4()}/sources/{s.id}/file');assert other.status_code==404
    path.unlink();assert (await client.get(base+'/file')).status_code==404

@pytest.mark.asyncio
@pytest.mark.parametrize('query',['kind=page&number=0','kind=page&number=-1','kind=page&number=3','kind=slide&number=1','kind=time&start=0','kind=page&number=1&path=../secret','kind=page&number=1&path=%2Fetc%2Fpasswd','kind=page&number=1&path=%2e%2e%2fsecret'])
async def test_invalid_locators_and_injection(db_session,client,query):
    c,s,_=await source(db_session)
    assert (await client.get(f'/courses/{c.id}/sources/{s.id}/location?'+query)).status_code==422
    assert (await client.get(f'/courses/{c.id}/sources/{s.id}/file?path=/etc/passwd')).status_code==422

@pytest.mark.asyncio
@pytest.mark.parametrize('query,status',[('kind=time&start=3&end=8',200),('kind=time&start=-1',422),('kind=time&start=nan',422),('kind=time&start=inf',422),('kind=time&start=11',422),('kind=time&start=3&end=2',422),('kind=time&start=3&end=11',422)])
async def test_video_bounds(db_session,client,monkeypatch,query,status):
    c,s,_=await source(db_session,'video')
    monkeypatch.setattr('app.services.source_preview.MediaTools.probe_media',lambda *a:{'format':{'duration':'10'}})
    result=await client.get(f'/courses/{c.id}/sources/{s.id}/location?'+query)
    assert result.status_code==status
    if status==200:assert result.json()['locator']['start']==3

@pytest.mark.asyncio
async def test_unknown_invalid_identity(client):
    assert (await client.get('/courses/not-uuid/sources/not-uuid/file')).status_code==422
    assert (await client.get(f'/courses/{uuid.uuid4()}/sources/{uuid.uuid4()}/file')).status_code==404

def test_symlink_and_absolute_path_rejected(tmp_path):
    root=tmp_path/'safe';root.mkdir();secret=tmp_path/'secret';secret.write_text('private')
    (root/'link').symlink_to(secret)
    for value in [str(secret),str(root/'link'),str(root/'../secret')]:
        with pytest.raises(HTTPException):contained_file(value,root)

@pytest.mark.asyncio
@pytest.mark.parametrize('supported',[True,False])
async def test_visual_real_asset_independent_verification(db_session,client,supported):
    c,s,_=await source(db_session)
    root=Path(settings.upload_dir).parent/'derived'/str(s.id);root.mkdir(parents=True)
    path=root/'diagram.png';Image.new('RGB',(20,20),'red').save(path)
    e=EvidenceUnit(id=uuid.uuid4(),course_id=c.id,source_version_id=s.id,modality='image',extraction_status='extracted',
        location_metadata={'page_number':2,'page_index':1,'derived_path':str(path)})
    db_session.add(e);await db_session.commit()
    class Provider:
        def __init__(self):self.calls=[]
        async def inspect(self,image,prompt):
            self.calls.append((image,prompt))
            return json.dumps({'evidence_type':'figure','description':'A red square.'} if len(self.calls)==1 else {'supported':supported,'reason':'Visible color checked.'})
    provider=Provider()
    result=await interpret(db_session,c.id,e.id,provider)
    await db_session.commit()
    assert len(provider.calls)==2 and all(v[0].startswith(b'\x89PNG') for v in provider.calls)
    assert result['semantic_state']==('verified' if supported else 'rejected')
    assert e.extracted_description==('A red square.' if supported else None)
    assert '/derived/' not in json.dumps(result,default=str)
    assert (await client.get(f'/courses/{c.id}/evidence/{e.id}/asset')).status_code==200
    assert (await client.get(f'/courses/{uuid.uuid4()}/evidence/{e.id}/asset')).status_code==404
    e.location_metadata={**e.location_metadata,'derived_path':'/etc/passwd'}
    with pytest.raises(HTTPException):visual_asset(e)

@pytest.mark.asyncio
async def test_authored_raster_only_gap_with_fixture_vision_then_grounded_ask(db_session,client,monkeypatch):
    """Actual frozen PDF raster -> vision fixture -> index/retrieval/Ask. Not a provider quality score."""
    import shutil
    from sqlalchemy import select
    from app.services.pdf_ingestion import process_pdf
    from app.services.indexing import index_evidence_units
    from app.services.grounded_ask import execute_grounded_ask
    from app.schemas.ask import AskRequest
    from app.models.evidence_unit import Modality
    from tests.test_api_assessment import SequentialFakeLLMProvider
    c,s,path=await source(db_session)
    authored=Path(__file__).resolve().parents[2]/'evaluation/datasets/sources/aster-memory.pdf'
    shutil.copyfile(authored,path)
    s.processing_status='uploaded';await db_session.commit()
    await process_pdf(db_session,source_version=s)
    rows=(await db_session.scalars(select(EvidenceUnit).where(EvidenceUnit.source_version_id==s.id,
        EvidenceUnit.modality=='image'))).all()
    raster=next(e for e in rows if e.location_metadata['page_number']==4)
    assert raster.original_text is None and raster.extracted_description is None
    original=visual_asset(raster)
    monkeypatch.setattr(settings,'vision_model',None)
    unavailable=await client.post(f'/courses/{c.id}/evidence/{raster.id}/interpret')
    assert unavailable.status_code==503
    detail=(await client.get(f'/courses/{c.id}/evidence/{raster.id}')).json()
    assert detail['asset_extracted'] and detail['semantic_state']=='not_interpreted'
    class FixtureVision:
        calls=0
        async def inspect(self,image,prompt):
            assert image==original
            self.calls+=1
            return json.dumps(dict(evidence_type='diagram',description='In the Aster routing figure, the dashed route from Amber gate terminates at Violet queue.') if self.calls==1 else dict(supported=True,reason='Fixture judgment for the visible raster labels and dashed route.'))
    await interpret(db_session,c.id,raster.id,FixtureVision());await db_session.commit()
    await index_evidence_units(db_session,c.id)
    assert raster.embedding is not None
    provider=SequentialFakeLLMProvider([json.dumps({'blocks':[{'text':'The dashed route terminates at Violet queue.','evidence_unit_ids':[str(raster.id)]}],'next_question':None})])
    answer=await execute_grounded_ask(db_session,c.id,AskRequest(question='Amber gate Violet queue',modalities=[Modality.image]),provider)
    assert answer.status=='answered'
    assert answer.blocks[0].citations[0].locator['page_number']==4
    assert answer.blocks[0].citations[0].modality==Modality.image
    assert 'dashed route' in provider.prompts[0]['user']

@pytest.mark.asyncio
async def test_actual_video_duration_and_byte_range(db_session,client):
    import shutil
    c,s,path=await source(db_session,'video')
    shutil.copyfile(Path(__file__).resolve().parents[2]/'evaluation/datasets/sources/aster-demonstration.mp4',path)
    base=f'/courses/{c.id}/sources/{s.id}'
    result=await client.get(base+'/location?kind=time&start=3&end=8')
    assert result.status_code==200,result.text
    assert 58 < result.json()['duration'] < 59
    result=await client.get(base+'/file',headers={'Range':'bytes=0-1023'})
    assert result.status_code==206
    assert result.headers['content-range']==f'bytes 0-1023/{path.stat().st_size}'
    assert result.content==path.read_bytes()[:1024]
